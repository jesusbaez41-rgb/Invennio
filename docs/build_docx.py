#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera la documentación integral del Repositorio HGI en formato DOCX.

Uso:
    /ruta/al/venv/bin/python docs/build_docx.py

Requiere: python-docx, markdown, beautifulsoup4 y lxml.
El contenido editorial vive en docs/doc_content.py; las matrices del estándar
MITIC se incluyen directamente desde docs/estandar-mitic/*.md.
"""

from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT,
                             WD_TAB_LEADER)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
OUT = HERE / "Documentacion_Repositorio_HGI.docx"

sys.path.insert(0, str(HERE))
from doc_content import CONTENT  # noqa: E402

# ---------------------------------------------------------------------------
# Identidad visual
# ---------------------------------------------------------------------------
BRAND = RGBColor(0x8C, 0x16, 0x18)
ACCENT = RGBColor(0xEA, 0x24, 0x28)
INK = RGBColor(0x22, 0x15, 0x18)
TEXT = RGBColor(0x3C, 0x32, 0x34)
MUTED = RGBColor(0x6F, 0x61, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEX_BRAND = "8C1618"
HEX_ACCENT = "EA2428"
HEX_SOFT = "FDF4F4"
HEX_LINE = "E3D6D7"
HEX_ZEBRA = "FDF9F9"

PAGE_W, PAGE_H = Cm(21.0), Cm(29.7)
MARGIN_LR, MARGIN_TOP, MARGIN_BOTTOM = Cm(2.5), Cm(2.5), Cm(2.2)
CONTENT_W = 16.0  # cm útiles en retrato
LAND_W = 25.7     # cm útiles en apaisado


# ---------------------------------------------------------------------------
# Utilidades XML
# ---------------------------------------------------------------------------
def _el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def shade(element_pr, fill: str) -> None:
    shd = _el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill})
    element_pr.append(shd)


def cell_shade(cell, fill: str) -> None:
    shade(cell._tc.get_or_add_tcPr(), fill)


def para_shade(paragraph, fill: str) -> None:
    shade(paragraph._p.get_or_add_pPr(), fill)


def para_border(paragraph, edge: str = "bottom", sz: int = 8,
                color: str = HEX_ACCENT, space: int = 4) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    b = _el("w:" + edge, **{"w:val": "single", "w:sz": sz,
                            "w:space": space, "w:color": color})
    pBdr.append(b)


def table_borders(table, color: str = HEX_LINE, sz: int = 4) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_el("w:" + edge, **{"w:val": "single", "w:sz": sz,
                                           "w:space": 0, "w:color": color}))
    tblPr.append(borders)


def table_fixed_layout(table) -> None:
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "pct")
    tblW.set(qn("w:w"), "5000")
    tblPr.append(tblW)


def cell_margins(table, top: int = 45, bottom: int = 45,
                 left: int = 80, right: int = 80) -> None:
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for name, val in (("top", top), ("left", left),
                      ("bottom", bottom), ("right", right)):
        mar.append(_el("w:" + name, **{"w:w": val, "w:type": "dxa"}))
    tblPr.append(mar)


def repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(_el("w:tblHeader", **{"w:val": "true"}))


def no_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(_el("w:cantSplit", **{"w:val": "true"}))


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = _el("w:fldChar", **{"w:fldCharType": "begin"})
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = _el("w:fldChar", **{"w:fldCharType": "separate"})
    text = OxmlElement("w:t")
    text.text = placeholder
    end = _el("w:fldChar", **{"w:fldCharType": "end"})
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(sep)
    run._r.append(text)
    run._r.append(end)


def add_hyperlink(paragraph, url: str, text: str) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), HEX_BRAND)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(color)
    rPr.append(underline)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


# ---------------------------------------------------------------------------
# Estilos
# ---------------------------------------------------------------------------
def setup_styles(doc: Document) -> None:
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Calibri")
    rfonts.set(qn("w:cs"), "Calibri")
    lang = _el("w:lang", **{"w:val": "es-PY", "w:eastAsia": "es-PY"})
    rpr.append(lang)

    h1 = styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(19)
    h1.font.bold = True
    h1.font.color.rgb = BRAND
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(14)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.page_break_before = True
    h1.paragraph_format.line_spacing = 1.0

    h2 = styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(13.5)
    h2.font.bold = True
    h2.font.color.rgb = BRAND
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.line_spacing = 1.0

    h3 = styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(11.5)
    h3.font.bold = True
    h3.font.color.rgb = INK
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True
    h3.paragraph_format.line_spacing = 1.0

    h4 = styles["Heading 4"]
    h4.font.name = "Calibri"
    h4.font.size = Pt(10.5)
    h4.font.bold = True
    h4.font.italic = False
    h4.font.color.rgb = RGBColor(0x5A, 0x2A, 0x2C)
    h4.paragraph_format.space_before = Pt(9)
    h4.paragraph_format.space_after = Pt(3)
    h4.paragraph_format.keep_with_next = True

    for name, indent in (("List Bullet", 0.55), ("List Number", 0.55)):
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(10.5)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.1
        st.paragraph_format.left_indent = Cm(indent)
        st.paragraph_format.first_line_indent = Cm(-0.35)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    code = styles.add_style("Código", 1)
    code.base_style = styles["Normal"]
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor(0x3A, 0x1A, 0x1B)
    code.paragraph_format.left_indent = Cm(0.35)
    code.paragraph_format.right_indent = Cm(0.2)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.keep_together = True

    quote = styles.add_style("Nota", 1)
    quote.base_style = styles["Normal"]
    quote.font.size = Pt(9.5)
    quote.font.italic = True
    quote.paragraph_format.left_indent = Cm(0.45)
    quote.paragraph_format.right_indent = Cm(0.3)
    quote.paragraph_format.space_before = Pt(4)
    quote.paragraph_format.space_after = Pt(6)
    quote.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    caption = styles.add_style("Pie de figura", 1)
    caption.base_style = styles["Normal"]
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = MUTED
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(10)

    tabla = styles.add_style("Tabla", 1)
    tabla.base_style = styles["Normal"]
    tabla.font.size = Pt(8)
    tabla.paragraph_format.space_before = Pt(0)
    tabla.paragraph_format.space_after = Pt(0)
    tabla.paragraph_format.line_spacing = 1.0
    tabla.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    tabla_h = styles.add_style("Tabla encabezado", 1)
    tabla_h.base_style = styles["Tabla"]
    tabla_h.font.bold = True
    tabla_h.font.color.rgb = WHITE


# ---------------------------------------------------------------------------
# Renderizador Markdown
# ---------------------------------------------------------------------------
EMOJI_MAP = [
    ("Plan de cierre de ítems 🟡", "Plan de cierre de ítems en curso"),
    ("✅ implementado", "implementado"),
    ("🟡 parcial", "parcial"),
    ("✅ ✅ ✅ ✅ ✅ ✅", "Cumple 6/6"),
    ("✅ Cumplido y verificado", "Cumplido y verificado"),
    ("✅ Cumplido", "Cumplido"),
    ("🟡 Parcial / en curso", "Parcial / en curso"),
    ("🟡 Parcial", "Parcial"),
    ("⚪ No aplica", "No aplica"),
    ("✅", "Cumplido"),
    ("🟡", "Parcial"),
    ("⚪", "No aplica"),
]


def normalize_text(text: str) -> str:
    for a, b in EMOJI_MAP:
        text = text.replace(a, b)
    return text


def shift_headings(text: str, drop_title: bool = True, levels: int = 1) -> str:
    lines = text.splitlines()
    if drop_title and lines and lines[0].startswith("# "):
        lines = lines[1:]
    while lines and not lines[0].strip():
        lines.pop(0)
    out, in_code = [], False
    for line in lines:
        if line.strip().startswith("```"):
            in_code = not in_code
        if not in_code and re.match(r"^#{1,6} ", line):
            line = "#" * levels + line
        out.append(line)
    return "\n".join(out)


def load_doc(rel_path: str, levels: int = 1) -> str:
    raw = (HERE / rel_path).read_text(encoding="utf-8")
    return shift_headings(raw, drop_title=True, levels=levels)


def expand_includes(text: str) -> str:
    """Sustituye {{DOC:ruta}} por el Markdown del archivo, jerarquizado."""
    def repl(match):
        return load_doc(match.group(1).strip(), levels=1)
    return re.sub(r"\{\{DOC:([^}]+)\}\}", repl, text)


def inline(par, node, bold=False, italic=False, code=False) -> None:
    """Vuelca nodos HTML en un párrafo conservando el formato básico."""
    if isinstance(node, NavigableString):
        text = str(node)
        if not code:
            text = re.sub(r"\s+", " ", text)
        if not text:
            return
        run = par.add_run(text)
        run.bold = bold
        run.italic = italic
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = BRAND
            rpr = run._r.get_or_add_rPr()
            shade(rpr, "F6EDEE")
        return

    if not isinstance(node, Tag):
        return

    name = node.name
    if name in ("strong", "b"):
        for c in node.children:
            inline(par, c, True, italic, code)
    elif name in ("em", "i"):
        for c in node.children:
            inline(par, c, bold, True, code)
    elif name in ("del", "s"):
        for c in node.children:
            inline(par, c, bold, italic, code)
    elif name == "code":
        for c in node.children:
            inline(par, c, bold, italic, True)
    elif name == "a":
        href = node.get("href", "")
        text = node.get_text()
        if href.startswith("http") and text.strip() and not any(
                isinstance(x, Tag) for x in node.children):
            add_hyperlink(par, href, text)
        else:
            for c in node.children:
                inline(par, c, bold, italic, code)
    elif name == "br":
        par.add_run().add_break()
    elif name == "img":
        return
    else:
        for c in node.children:
            inline(par, c, bold, italic, code)


def add_paragraph(doc, node, style=None, base_size=None, align=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    inline(p, node)
    if base_size:
        for r in p.runs:
            r.font.size = Pt(base_size)
    if align is not None:
        p.alignment = align
    return p


CELL_BLOCKS = {"p", "div", "ul", "ol", "pre", "blockquote", "table",
               "h1", "h2", "h3", "h4", "h5", "h6", "hr"}


def render_table(doc, table_tag, width_cm: float):
    rows = []
    header = None
    thead = table_tag.find("thead")
    if thead:
        header = [th for th in thead.find_all("tr")[0].find_all(["th", "td"])]
    body_rows = table_tag.find("tbody")
    trs = body_rows.find_all("tr", recursive=False) if body_rows else table_tag.find_all("tr")
    for tr in trs:
        cells = tr.find_all(["td", "th"], recursive=False)
        if cells:
            rows.append(cells)
    if header:
        rows.insert(0, header)

    if not rows:
        return

    ncols = max(len(r) for r in rows)
    # Ancho proporcional al contenido, con mínimos por columna que respetan
    # la palabra más larga (evita cortes como «DATABAS / E_URI»).
    stats = []
    for c in range(ncols):
        lens, longest = [], 0
        for r in rows:
            if c < len(r):
                txt = r[c].get_text(" ", strip=True)
                lens.append(max(len(txt), 4))
                for token in re.split(r"[\s/]+", txt):
                    longest = max(longest, len(token))
        avg = sum(lens) / len(lens) if lens else 10
        stats.append((max(avg, 8), longest))

    weights = [avg ** 0.75 for avg, _ in stats]
    total = sum(weights)
    widths = [width_cm * w / total for w in weights]
    floors = [min(max(1.45, 0.165 * longest + 0.25), width_cm * 0.42)
              for _, longest in stats]

    for _ in range(8):
        deficit = sum(max(0.0, floors[i] - widths[i]) for i in range(ncols))
        if deficit < 0.05:
            break
        free = [i for i in range(ncols) if widths[i] > floors[i]]
        if not free:
            break
        pool = sum(widths[i] - floors[i] for i in free)
        for i in range(ncols):
            if widths[i] < floors[i]:
                widths[i] = floors[i]
        for i in free:
            widths[i] -= deficit * (widths[i] - floors[i]) / pool

    scale = width_cm / sum(widths)
    widths = [w * scale for w in widths]
    dense = max(longest for _, longest in stats) > 26
    table_size = Pt(7.5) if dense else Pt(8)

    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_fixed_layout(table)
    table_borders(table)
    cell_margins(table)

    for ri, cells in enumerate(rows):
        row = table.rows[ri]
        is_header = (ri == 0 and header is not None)
        if is_header:
            repeat_header(row)
        no_split(row)
        for ci in range(ncols):
            cell = row.cells[ci]
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if is_header:
                cell_shade(cell, HEX_BRAND)
            elif ri % 2 == 0:
                cell_shade(cell, HEX_ZEBRA)

            # limpiar párrafo inicial
            first = cell.paragraphs[0]
            first.style = "Tabla"
            for r in list(first.runs):
                r._r.getparent().remove(r._r)
            if ci < len(cells):
                node = cells[ci]
                blocks = [c for c in node.children
                          if isinstance(c, Tag) and c.name in CELL_BLOCKS]
                if not blocks:
                    inline(first, node)
                    if is_header:
                        first.style = "Tabla encabezado"
                else:
                    target = first
                    for bi, block in enumerate(blocks):
                        if block.name in ("p", "div"):
                            if bi > 0:
                                target = cell.add_paragraph(style="Tabla")
                            inline(target, block)
                        elif block.name in ("ul", "ol"):
                            for li in block.find_all("li", recursive=False):
                                target = cell.add_paragraph(style="Tabla")
                                target.paragraph_format.left_indent = Cm(0.25)
                                target.add_run("• ").bold = True
                                inline(target, li)
                        elif block.name == "pre":
                            target = cell.add_paragraph(style="Tabla")
                            target.add_run(block.get_text()).font.name = "Consolas"
                        else:
                            if bi > 0:
                                target = cell.add_paragraph(style="Tabla")
                            inline(target, block)
                if is_header:
                    for p in cell.paragraphs:
                        p.style = "Tabla encabezado"
                        for r in p.runs:
                            r.font.color.rgb = WHITE
                            r.font.bold = True

    if dense:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = table_size

    # espaciado posterior a la tabla
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 0.6
    for r in spacer.runs:
        r.font.size = Pt(2)


def render_list(doc, node, ordered=False, level=0):
    """Listas con numeración explícita: cada lista reinicia en 1 y el número
    no depende de la posición que Word o LibreOffice asignen al estilo."""
    for idx, li in enumerate(node.find_all("li", recursive=False), start=1):
        # separar el contenido propio del li de las sublistas
        sublists = [c for c in li.children if isinstance(c, Tag)
                    and c.name in ("ul", "ol")]
        if ordered:
            p = doc.add_paragraph()
            base = 0.9 + 0.55 * level
            pf = p.paragraph_format
            pf.left_indent = Cm(base)
            pf.first_line_indent = Cm(-0.55)
            pf.space_before = Pt(0)
            pf.space_after = Pt(3)
            pf.line_spacing = 1.12
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.tab_stops.add_tab_stop(Cm(base), WD_TAB_ALIGNMENT.LEFT)
            num = p.add_run(f"{idx}.")
            num.bold = True
            num.font.color.rgb = BRAND
            p.add_run("\t")
        else:
            p = doc.add_paragraph(style="List Bullet")
            if level:
                p.paragraph_format.left_indent = Cm(0.55 + 0.5 * level)
        for c in li.children:
            if isinstance(c, Tag) and c.name in ("ul", "ol"):
                continue
            inline(p, c)
        for sub in sublists:
            render_list(doc, sub, sub.name == "ol", level + 1)


def render_blockquote(doc, node):
    for child in node.children:
        if isinstance(child, Tag) and child.name in ("p", "div"):
            p = add_paragraph(doc, child, style="Nota")
            para_border(p, "left", sz=18, color=HEX_ACCENT, space=8)
            para_shade(p, HEX_SOFT)
        elif isinstance(child, Tag) and child.name in ("ul", "ol"):
            render_list(doc, child, child.name == "ol")


def render_code(doc, node):
    code_node = node.find("code") or node
    text = code_node.get_text().rstrip("\n")
    p = doc.add_paragraph(style="Código")
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        p.add_run(line if line else " ")
    para_border(p, "left", sz=18, color=HEX_ACCENT, space=6)
    para_shade(p, "F8F4F4")


def render_markdown(doc, text: str, base_dir: Path, width_cm: float):
    html = markdown.markdown(
        text,
        extensions=["tables", "fenced_code", "sane_lists"],
        output_format="html5",
    )
    soup = BeautifulSoup(html, "lxml")
    body = soup.body or soup
    for node in body.children:
        if isinstance(node, NavigableString):
            if node.strip():
                doc.add_paragraph(node.strip())
            continue
        if not isinstance(node, Tag):
            continue
        name = node.name
        if name == "h1":
            p = doc.add_heading(level=1)
            inline(p, node)
        elif name == "h2":
            p = doc.add_heading(level=2)
            inline(p, node)
        elif name == "h3":
            p = doc.add_heading(level=3)
            inline(p, node)
        elif name == "h4":
            p = doc.add_heading(level=4)
            inline(p, node)
        elif name in ("h5", "h6"):
            p = doc.add_heading(level=4)
            inline(p, node)
        elif name == "p":
            p = add_paragraph(doc, node)
        elif name == "ul":
            render_list(doc, node, False)
        elif name == "ol":
            render_list(doc, node, True)
        elif name == "table":
            render_table(doc, node, width_cm)
        elif name == "pre":
            render_code(doc, node)
        elif name == "blockquote":
            render_blockquote(doc, node)
        elif name == "hr":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            para_border(p, "bottom", sz=6, color=HEX_LINE, space=1)
        elif name == "img":
            src = node.get("src", "")
            path = (base_dir / src).resolve()
            if path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(path), width=Cm(min(width_cm, 15)))
        else:
            inline(doc.add_paragraph(), node)


TOKEN_RE = re.compile(
    r"^<<(PAGEBREAK|TOC|LANDSCAPE|PORTRAIT|FIG:[^>]+)>>[ \t]*$", re.MULTILINE)

TOC_ENTRIES: list[tuple[int, str]] = []
PAGE_MAP: dict[str, int] = {}


def add_rich(doc, text: str, base_dir: Path | None = None) -> None:
    base_dir = base_dir or HERE
    text = normalize_text(text)
    pos = 0
    for m in TOKEN_RE.finditer(text):
        chunk = text[pos:m.start()]
        if chunk.strip():
            render_markdown(doc, chunk, base_dir, current_width(doc))
        token = m.group(1)
        if token == "PAGEBREAK":
            doc.add_page_break()
        elif token == "TOC":
            add_toc(doc, TOC_ENTRIES, PAGE_MAP)
        elif token == "LANDSCAPE":
            new_section(doc, landscape=True)
        elif token == "PORTRAIT":
            new_section(doc, landscape=False)
        elif token.startswith("FIG:"):
            payload = token[4:].split("|")
            src = payload[0].strip()
            caption = payload[1].strip() if len(payload) > 1 else ""
            width = float(payload[2]) if len(payload) > 2 else 15.0
            path = (base_dir / src).resolve()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=Cm(width))
            if caption:
                c = doc.add_paragraph(style="Pie de figura")
                c.add_run(caption)
        pos = m.end()
    chunk = text[pos:]
    if chunk.strip():
        render_markdown(doc, chunk, base_dir, current_width(doc))


def current_width(doc) -> float:
    s = doc.sections[-1]
    if s.orientation == WD_ORIENT.LANDSCAPE:
        return LAND_W
    return CONTENT_W


def new_section(doc, landscape: bool):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        s.page_width, s.page_height = PAGE_H, PAGE_W
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin, s.bottom_margin = Cm(2.2), Cm(2.0)
    else:
        s.page_width, s.page_height = PAGE_W, PAGE_H
        s.left_margin = s.right_margin = MARGIN_LR
        s.top_margin, s.bottom_margin = MARGIN_TOP, MARGIN_BOTTOM
    s.header_distance = Cm(1.1)
    s.footer_distance = Cm(1.0)
    s.different_first_page_header_footer = False
    s.gutter = Cm(0)
    return s


# ---------------------------------------------------------------------------
# Portada, encabezados y pie
# ---------------------------------------------------------------------------
def add_cover(doc: Document) -> None:
    # logo
    logo = HERE / "assets/documentacion/logo-mspbs.jpg"
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(logo), width=Cm(8.6))

    line = doc.add_paragraph()
    line.paragraph_format.space_before = Pt(0)
    line.paragraph_format.space_after = Pt(42)
    para_border(line, "bottom", sz=18, color=HEX_ACCENT, space=2)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(0)
    kicker.paragraph_format.space_after = Pt(8)
    r = kicker.add_run("DOCUMENTACIÓN TÉCNICA DEL SISTEMA")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    r.font.name = "Calibri"

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.line_spacing = 0.95
    r = title.add_run("Repositorio HGI")
    r.bold = True
    r.font.size = Pt(44)
    r.font.color.rgb = BRAND

    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(16)
    r = sub.add_run("Repositorio Institucional de Ciencia Abierta")
    r.font.size = Pt(18)
    r.font.color.rgb = INK

    desc = doc.add_paragraph()
    desc.paragraph_format.space_after = Pt(34)
    desc.paragraph_format.line_spacing = 1.25
    r = desc.add_run(
        "Documentación integral del sistema: arquitectura, funcionalidades, "
        "operación, seguridad y cumplimiento del Estándar de Software del "
        "Ministerio de Tecnologías de la Información y Comunicación (MITIC).")
    r.font.size = Pt(11.5)
    r.font.color.rgb = TEXT

    # bloque institucional
    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = band.rows[0].cells[0]
    cell.width = Cm(16)
    cell_shade(cell, HEX_BRAND)
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(10)
    cp.paragraph_format.space_after = Pt(10)
    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = cp.add_run("  Hospital General de Itapúa · Encarnación, Paraguay")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = WHITE

    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(26)

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    meta_data = [
        ("Versión del documento", "1.0"),
        ("Fecha de emisión", "27 de octubre de 2026"),
        ("Clasificación", "Documento técnico de uso institucional"),
        ("Elaborado por", "Equipo de desarrollo del proyecto · Unidad TIC"),
    ]
    for i, (k, v) in enumerate(meta_data):
        c0, c1 = meta.rows[i].cells
        c0.width = Cm(5.4)
        c1.width = Cm(10.6)
        p0 = c0.paragraphs[0]
        p1 = c1.paragraphs[0]
        r0 = p0.add_run(k.upper())
        r0.bold = True
        r0.font.size = Pt(8)
        r0.font.color.rgb = MUTED
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = INK
        p0.paragraph_format.space_after = Pt(4)
        p1.paragraph_format.space_after = Pt(4)

    foot = doc.add_paragraph()
    foot.paragraph_format.space_before = Pt(30)
    foot.paragraph_format.space_after = Pt(0)
    para_border(foot, "top", sz=6, color=HEX_LINE, space=6)
    r = foot.add_run("Basado en InvenioRDM 13.0.8  ·  Norma de referencia: Estándar de Software — MITIC v1.0")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED


def add_header_footer(doc: Document) -> None:
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True

    header = sec.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run("Repositorio HGI  ·  Documentación integral del sistema")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    r.font.name = "Calibri"
    para_border(hp, "bottom", sz=6, color=HEX_LINE, space=4)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Ministerio de Salud Pública y Bienestar Social · Hospital General de Itapúa  ·  Página ")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    add_field(fp, "PAGE", "1")
    r = fp.add_run(" de ")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    add_field(fp, "NUMPAGES", "1")
    para_border(fp, "top", sz=6, color=HEX_LINE, space=4)


def add_toc(doc: Document, entries, page_map) -> None:
    """Índice estático con puntos guía y numeración de esta edición."""
    for level, title in entries:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.0 if level == 1 else 0.45 * (level - 1))
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(6 if level == 1 else 1)
        pf.space_after = Pt(1 if level == 1 else 0)
        pf.line_spacing = 1.0
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.tab_stops.add_tab_stop(
            Cm(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = p.add_run(title)
        run.font.size = Pt(10.5 if level == 1 else 9.5)
        run.bold = (level == 1)
        run.font.color.rgb = BRAND if level == 1 else (INK if level == 2 else MUTED)
        page = page_map.get(re.sub(r"\s+", " ", title).strip(), "")
        num = p.add_run("\t" + (str(page) if page else "—"))
        num.font.size = Pt(9.5 if level > 1 else 10)
        num.font.color.rgb = MUTED
    note = doc.add_paragraph(style="Pie de figura")
    note.add_run("Los números de página corresponden a esta edición del documento.")


def collect_toc_entries(text: str):
    entries, in_code = [], False
    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if not m:
            continue
        title = m.group(2)
        title = re.sub(r"`([^`]*)`", r"\1", title)
        title = title.replace("**", "").replace("__", "")
        title = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", title)
        title = title.replace("*", "").replace("&amp;", "&")
        title = re.sub(r"<[^>]+>", "", title)
        title = normalize_text(title)
        title = title.strip()
        entries.append((len(m.group(1)), title))
    return entries


# ---------------------------------------------------------------------------
# Construcción
# ---------------------------------------------------------------------------
def build() -> Path:
    global TOC_ENTRIES, PAGE_MAP

    doc = Document()
    setup_styles(doc)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = PAGE_W, PAGE_H
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.left_margin = sec.right_margin = MARGIN_LR
    sec.top_margin, sec.bottom_margin = MARGIN_TOP, MARGIN_BOTTOM
    sec.header_distance = Cm(1.1)
    sec.footer_distance = Cm(1.0)
    sec.gutter = Cm(0)

    add_header_footer(doc)

    props = doc.core_properties
    props.title = "Repositorio HGI — Documentación integral del sistema"
    props.subject = "Documentación técnica y de cumplimiento (Estándar de Software MITIC)"
    props.author = "Equipo de desarrollo del Repositorio HGI"
    props.category = "Documentación de sistema"
    props.comments = ("Documento consolidado: arquitectura, módulos, operación, "
                      "seguridad y matrices de cumplimiento MITIC.")
    props.keywords = "InvenioRDM; repositorio institucional; MITIC; HGI; documentación"
    props.language = "es-PY"

    add_cover(doc)
    expanded = expand_includes(CONTENT)
    TOC_ENTRIES = collect_toc_entries(expanded)
    map_path = os.environ.get("TOC_PAGES", "")
    if map_path and Path(map_path).exists():
        PAGE_MAP = json.loads(Path(map_path).read_text(encoding="utf-8"))
    add_rich(doc, expanded)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Documento generado: {path}")
